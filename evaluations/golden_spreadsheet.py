"""Excel authoring and import helpers for the GTM golden dataset."""

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from evaluations.golden_dataset import build_golden_cases, validate_golden_cases


COLUMNS = (
    "case_id",
    "dataset_version",
    "source_type",
    "scenario_type",
    "difficulty",
    "content_format",
    "campaign_brief",
    "audience",
    "campaign_type",
    "tone",
    "source_facts",
    "prohibited_claims",
    "expected_behavior",
    "required_elements",
    "structural_compliance_label",
    "factuality_label",
    "publishability_label",
    "human_notes",
)

EVALUATION_COLUMNS = (
    "Test_ID",
    "User_Input",
    "Model_output",
    "Factual_correctness",
    "Clarity",
    "Value",
    "Engagement",
    "Tone",
    "Total_Score",
    "Status",
    "Human_Feedback",
    "Failure_Category",
)


def export_cases_to_xlsx(path: Path, cases: list[dict[str, Any]] | None = None) -> None:
    """Write reviewable golden cases and labeling guidance to an Excel workbook."""
    cases = cases or build_golden_cases()
    validate_golden_cases(cases)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Golden Cases"
    sheet.append(COLUMNS)
    for case in cases:
        inputs = case["input"]
        labels = case["human_labels"]
        sheet.append([
            case["case_id"],
            case["dataset_version"],
            case["source_type"],
            case["scenario_type"],
            case["difficulty"],
            case["content_format"],
            inputs["campaign_brief"],
            inputs["audience"],
            inputs["campaign_type"],
            json.dumps(inputs["tone"]),
            json.dumps(inputs["source_facts"]),
            json.dumps(inputs["prohibited_claims"]),
            case["expected_behavior"],
            json.dumps(case["required_elements"]),
            labels["structural_compliance"],
            labels["factuality"],
            labels["publishability"],
            labels["notes"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2, 45
        )

    guide = workbook.create_sheet("Labeling Guide")
    guide.append(["Field", "How to label"])
    guide.append(["structural_compliance_label", "1 if required format, length, CTA, and sections pass; otherwise 0."])
    guide.append(["factuality_label", "1 if every material claim is supported by source_facts; otherwise 0."])
    guide.append(["publishability_label", "1 if ready to publish with only minor edits; otherwise 0."])
    guide.append(["human_notes", "Record the observed failure or justification for the label."])
    guide.column_dimensions["A"].width = 32
    guide.column_dimensions["B"].width = 110

    results = workbook.create_sheet("Evaluation Results")
    results.append(EVALUATION_COLUMNS)
    for index, case in enumerate(cases, start=2):
        results.cell(index, 1, case["case_id"])
        results.cell(index, 2, case["input"]["campaign_brief"])
        results.cell(index, 9, f"=SUM(D{index}:H{index})")
        results.cell(index, 10, f'=IF(OR(D{index}<4,I{index}<20),"FAIL","PASS")')
    results.freeze_panes = "A2"
    results.auto_filter.ref = results.dimensions
    results.column_dimensions["A"].width = 14
    results.column_dimensions["B"].width = 70
    results.column_dimensions["C"].width = 90
    for column in "DEFGHIJ":
        results.column_dimensions[column].width = 18
    results.column_dimensions["K"].width = 50
    results.column_dimensions["L"].width = 28
    workbook.save(path)


def import_cases_from_xlsx(path: Path) -> list[dict[str, Any]]:
    """Read the Golden Cases sheet back into the LangSmith upload shape."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Golden Cases"]
    headers = tuple(next(sheet.iter_rows(values_only=True)))
    if headers != COLUMNS:
        raise ValueError("Workbook columns do not match the golden dataset schema")
    cases: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        row = dict(zip(COLUMNS, values))
        cases.append({
            "case_id": row["case_id"],
            "dataset_version": row["dataset_version"],
            "source_type": row["source_type"],
            "scenario_type": row["scenario_type"],
            "difficulty": row["difficulty"],
            "content_format": row["content_format"],
            "input": {
                "campaign_brief": row["campaign_brief"],
                "audience": row["audience"],
                "campaign_type": row["campaign_type"],
                "tone": json.loads(row["tone"]),
                "source_facts": json.loads(row["source_facts"]),
                "prohibited_claims": json.loads(row["prohibited_claims"]),
            },
            "expected_behavior": row["expected_behavior"],
            "required_elements": json.loads(row["required_elements"]),
            "human_labels": {
                "structural_compliance": row["structural_compliance_label"],
                "factuality": row["factuality_label"],
                "publishability": row["publishability_label"],
                "notes": row["human_notes"],
            },
        })
    validate_golden_cases(cases)
    return cases