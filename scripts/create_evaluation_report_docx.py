#!/usr/bin/env python3
"""Convert the GTM evaluation Markdown report into a Word document."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "GTM_EVALUATION_REPORT.md"
OUTPUT = ROOT / "docs" / "GTM_EVALUATION_REPORT.docx"


def add_table(document: Document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(set(cell) <= {"-", ":", " "} for cell in row)]
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, rows[0]):
        cell.text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.add_paragraph()


def build_document() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            paragraph.style = "No Spacing"
            run = paragraph.add_run("\n".join(code))
            run.font.name = "Courier New"
            run.font.size = Pt(8)
            index += 1
            continue
        if line.startswith("# "):
            heading = document.add_heading(line[2:].strip(), level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            index += 1
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        if line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
            index += 1
            continue
        document.add_paragraph(line)
        index += 1

    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_document()