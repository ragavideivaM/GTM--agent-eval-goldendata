#!/usr/bin/env python3
"""Create the submission-ready AI at Your Doorstep project report."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/AI_at_Your_Doorstep_Assignment_Report.docx")
BLUE = "245B78"
DARK = "183446"
GREEN = "3E8E66"
LIGHT_BLUE = "EAF3F7"
LIGHT_GREEN = "EAF5EF"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
BLACK = "1F2937"
TABLE_WIDTH = 9360


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_run(run, size=11, bold=False, color=BLACK, italic=False, font="Calibri") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, after=6, before=0, line=1.1) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_body(doc, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    style_paragraph(p)
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, after=5, line=1.1)
        set_run(p.add_run(item))


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        style_paragraph(p, after=5, line=1.1)
        set_run(p.add_run(item))


def add_heading(doc, text: str, level=1, *, page_break_before=False) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before


def add_callout(doc, label: str, text: str, fill=LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    style_paragraph(p, after=0, line=1.1)
    set_run(p.add_run(f"{label}: "), bold=True, color=DARK)
    set_run(p.add_run(text), color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, BLUE)
        p = cell.paragraphs[0]
        style_paragraph(p, after=0)
        set_run(p.add_run(header), size=9.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                shade(cells[index], "F8FAFC")
            p = cells[index].paragraphs[0]
            style_paragraph(p, after=0, line=1.05)
            set_run(p.add_run(value), size=9.3)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_flow(doc, steps: list[str]) -> None:
    table = doc.add_table(rows=1, cols=len(steps))
    widths = [TABLE_WIDTH // len(steps)] * len(steps)
    widths[-1] += TABLE_WIDTH - sum(widths)
    set_table_geometry(table, widths)
    for index, step in enumerate(steps):
        cell = table.cell(0, index)
        shade(cell, LIGHT_GREEN if index % 2 else LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        set_run(p.add_run(step), size=9.2, bold=True, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run(run, size=9, color=MID_GRAY)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run("AI AT YOUR DOORSTEP  |  GTM AGENT PROJECT REPORT"), size=8.5, bold=True, color=MID_GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(footer.add_run("Assignment submission  •  "), size=9, color=MID_GRAY)
    add_page_number(footer)


def cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(108)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("PROJECT REPORT"), size=11, bold=True, color=GREEN)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=8)
    set_run(p.add_run("AI at Your Doorstep"), size=30, bold=True, color=DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=24)
    set_run(p.add_run("A Human-Approved, RAG-Grounded Go-to-Market Content Agent"), size=15, color=BLUE)
    add_callout(
        doc,
        "Project outcome",
        "A working MVP that discovers weekly AI stories, grounds claims in approved evidence, generates four GTM formats, runs automated review, and preserves human approval before publication.",
        LIGHT_GREEN,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Prepared as an assignment submission"), size=11, bold=True, color=MID_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("August 2026"), size=10.5, color=MID_GRAY)


def build() -> Document:
    doc = Document()
    configure_document(doc)
    cover(doc)

    add_heading(doc, "Executive summary", page_break_before=True)
    add_body(doc, "AI at Your Doorstep is a human-approved go-to-market agent for producing a weekly “What’s New in AI?” campaign. It discovers timely AI stories, lets a user select trusted sources, converts those sources into retrievable evidence, and produces a LinkedIn post, promotional email, short blog draft, and ad variations. A separate review agent critiques and revises the suite before final human approval.")
    add_body(doc, "The MVP combines a Streamlit interface, OpenAI models and embeddings, Pinecone semantic retrieval, Pydantic structured outputs, and an optional Deep Agent coordinator. Its central design principle is bounded autonomy: models perform research synthesis, writing, and review, while people retain control over source selection, paid API operations, brief approval, revision feedback, final approval, and publishing.")
    add_callout(doc, "Verified result", "The final implementation has 71 passing automated tests. A ten-document representative corpus and 25-question retrieval challenge supported the selected 3,600-character chunk size with 450-character overlap.")

    add_heading(doc, "1. Problem definition and objectives")
    add_heading(doc, "1.1 Problem", 2)
    add_body(doc, "Busy professionals need a concise view of meaningful AI developments, but weekly research and multi-channel campaign production are repetitive and time-consuming. A useful system must move faster without sacrificing source quality, factual grounding, tone consistency, or human editorial control.")
    add_heading(doc, "1.2 Primary objective", 2)
    add_body(doc, "Build awareness for the AI Week in Review newsletter and support the first 100 subscribers by generating credible, practical, ready-to-edit GTM content from current or uploaded evidence.")
    add_heading(doc, "1.3 MVP requirements", 2)
    add_bullets(doc, [
        "Accept newsletter, feature, or upcoming-event campaigns.",
        "Discover current AI stories and accept PDF, Excel, or pasted-text inputs.",
        "Index only sources explicitly approved by a human.",
        "Use RAG to ground a structured weekly research brief.",
        "Generate LinkedIn, email, blog, and advertising assets.",
        "Run an automated reviewer and support human revision feedback.",
        "Require final human approval and never publish automatically.",
    ])

    add_heading(doc, "2. Solution architecture")
    add_flow(doc, ["Discover / upload", "Human approval", "Extract + chunk", "Embed + index", "Retrieve + brief", "Generate + review", "Final approval"])
    add_body(doc, "The application separates deterministic workflow controls from model-driven work. Streamlit owns session state and approval gates. The research workflow owns discovery, ingestion, retrieval, and brief construction. The agent layer owns structured writing and review. Pinecone stores vectors and source metadata, while Pydantic contracts enforce predictable data at every boundary.")
    add_table(doc, ["Layer", "Technology", "Responsibility"], [
        ["Experience", "Streamlit", "Four-tab interface, session state, approvals, cost notices, exports"],
        ["Research", "OpenAI web search + local parsers", "Discover, extract, normalize, and prepare trusted evidence"],
        ["Retrieval", "OpenAI embeddings + Pinecone", "Semantic indexing, filtering, ranking, and evidence return"],
        ["Generation", "OpenAI Responses API + Pydantic", "Grounded brief and four structured content formats"],
        ["Orchestration", "Deep Agents + approval-aware tools", "Bounded generation/review sequence with fallback"],
        ["Quality", "Reviewer + deterministic evaluations", "Grounding references, tone, consistency, CTA, regression checks"],
    ], [1500, 2500, 5360])

    add_heading(doc, "3. End-to-end workflow")
    add_numbered(doc, [
        "Choose the campaign type and optionally upload a PDF or Excel source.",
        "Discover recent AI stories for the automatically calculated seven-day window.",
        "Review candidates and approve only trusted sources.",
        "Extract, normalize, chunk, embed, and index approved sources in Pinecone.",
        "Run multi-query retrieval and generate a structured, evidence-citing research brief.",
        "Review the brief and approve the combined paid operation.",
        "Generate four content formats and run the review agent in one bounded workflow.",
        "Apply human feedback if required, then provide final approval and export JSON.",
    ])
    add_callout(doc, "Human-in-the-loop boundary", "Discovery results are suggestions, not evidence. A source becomes usable only after explicit approval and successful indexing. Final approval never triggers automatic publication.")

    add_heading(doc, "4. RAG ingestion and retrieval design")
    add_heading(doc, "4.1 Ingestion", 2)
    add_body(doc, "The ingestion layer supports web pages, text-based PDFs, Excel workbooks, and pasted text. Parsers normalize content and preserve useful metadata such as title, publisher, source type, publication date, week identifier, company, topic, and source URL. Stable document and chunk IDs make citations traceable across Pinecone, the research brief, generated assets, and evaluation bundles.")
    add_heading(doc, "4.2 Corpus-tied chunking", 2)
    add_body(doc, "Chunking was treated as an upstream design choice rather than an arbitrary constant. The implementation prefers paragraph boundaries, splits oversized paragraphs at word boundaries, and adds a word-aligned overlap tail. Ten representative documents covered two industry-news articles, two tool announcements, two events, two product features, and two research documents.")
    add_table(doc, ["Configuration", "Source recall @8", "Strict passage recall @8", "Decision"], [
        ["1,200 / 150", "100.0%", "28.3%", "Rejected: answer-bearing context frequently split"],
        ["2,400 / 300", "84.8%", "23.2%", "Rejected: weakest challenge performance"],
        ["3,600 / 450", "100.0%", "72.5%", "Selected production default"],
    ], [1700, 1900, 2100, 3660])
    add_body(doc, "The comparison used 25 questions, source-deduplicated results, top-k values of 5, 8, and 10, and three repeated runs. Top-k 8 was retained for the evaluation because top-k 10 produced no measurable benefit. No-answer similarity scores were reported separately because their ranges overlapped and did not support a dependable abstention threshold.")
    add_heading(doc, "4.3 Retrieval", 2)
    add_body(doc, "The production retriever embeds several editorial queries covering launches, research, policy and safety, practical tools, and major industry developments. Pinecone queries are filtered by week, results are deduplicated by chunk ID, and the highest-scoring evidence is passed to the brief builder. The current MVP is semantic-only; a lexical or hybrid BM25 branch remains future work.")

    add_heading(doc, "5. Agent and content-generation design")
    add_flow(doc, ["Approved brief", "Deep Agent", "Writer tool", "Structured suite", "Reviewer tool", "Revised suite"])
    add_heading(doc, "5.1 Main coordinator", 2)
    add_body(doc, "GTMMainAgent exposes a small deterministic interface for discovery, indexing, brief creation, content generation, and review. This prevents orchestration logic from bypassing approval rules and makes the same capabilities testable without the UI.")
    add_heading(doc, "5.2 Deep Agent mode", 2)
    add_body(doc, "The Deep Agent coordinates only approved tools. It checks workflow state, reads the approved brief, invokes content generation once, and invokes review once. Middleware limits orchestration model calls, while the graph recursion limit allows internal planning and tool transitions. If the model skips an approved step, a deterministic fallback can complete that step; a failed paid call is not silently repeated.")
    add_heading(doc, "5.3 Structured content and review", 2)
    add_body(doc, "The writer returns a Pydantic ContentSuite containing a 100–350 word LinkedIn post, an email with exactly three subject lines, a 500–1,200 word blog draft, and three to five advertisements. Every factual claim carries evidence chunk IDs. The reviewer scores factual grounding, cross-format consistency, tone alignment, clarity, and CTA quality, then returns an approved or revised suite.")

    add_heading(doc, "6. Safety, governance, and cost controls")
    add_table(doc, ["Control", "Implementation"], [
        ["Source approval", "Only explicitly selected candidates are indexed."],
        ["Cost approval", "The UI displays conservative ceilings and requires approval before paid stages."],
        ["Brief approval", "Content tools remain blocked until a human approves the research brief."],
        ["Grounding", "Unknown evidence IDs are rejected at brief, generation, review, and bundle boundaries."],
        ["Bounded calls", "Generation and review are single-use per agent run; orchestration calls are limited."],
        ["Publishing", "No social, email, or advertising platform is connected for automatic publishing."],
        ["Secrets", ".env is excluded from version control; API keys are not stored in fixtures or source code."],
    ], [2300, 7060])
    add_body(doc, "Cost notices distinguish actual recorded usage from conservative ceilings. OpenAI charges are billed to the API project associated with the key in .env; Pinecone usage is billed according to the configured Pinecone account and plan. Displayed estimates are not provider-enforced spending caps.")

    add_heading(doc, "7. Evaluation strategy and results")
    add_heading(doc, "7.1 Deterministic quality gates", 2)
    add_bullets(doc, [
        "Schema validation for briefs, content suites, reviews, and evaluation bundles.",
        "LinkedIn and blog length, email subject-line count, and advertisement count.",
        "Campaign-appropriate CTA coverage in every reader-facing format.",
        "Non-empty evidence references and shared evidence across the four formats.",
        "Known risky marketing phrases and evidence-reference integrity.",
        "Golden passing fixture and intentionally flawed regression fixture.",
    ])
    add_heading(doc, "7.2 Retrieval evaluation", 2)
    add_table(doc, ["Evaluation", "Dataset", "Outcome"], [
        ["Baseline retrieval", "10 labelled corpus questions", "100% Hit@5; MRR 1.000"],
        ["Challenge retrieval", "25 paraphrased, multi-source, and no-answer questions", "Selected 3,600 / 450; 100% source recall @8"],
        ["Strict passage check", "Required answer terms in returned chunks", "72.5% passage recall @8 for selected configuration"],
        ["Automated tests", "Unit and regression suite", "71 tests passed"],
        ["Final suite export", "LinkedIn, email, blog, four ads", "All deterministic content checks passed"],
    ], [1900, 3500, 3960])
    add_callout(doc, "Interpretation", "Evidence-ID checks prove referential integrity: a citation exists and is known. They do not by themselves prove semantic entailment of every generated sentence. That stronger evaluation is explicitly retained as future work.")

    add_heading(doc, "8. User interface and demonstration")
    add_body(doc, "The Streamlit experience is organized into Welcome, Generate, Full Goto Content, and Debug tabs. The Welcome tab introduces the application and accepts campaign type and optional documents. Generate presents numbered stages and approval gates. Full Goto Content contains the brief, review report, revised assets, feedback interface, final approval, and export. Debug exposes structured state needed during development.")
    add_heading(doc, "Recommended demonstration sequence", 2)
    add_numbered(doc, [
        "Start Streamlit and open the local URL printed in the terminal.",
        "Discover five current stories, approve at least three, and index them.",
        "Show the source count, chunk count, embedding usage, and estimated charge.",
        "Generate the grounded brief and inspect source links and evidence IDs.",
        "Approve the combined operation and run the agent workflow.",
        "Open Full Goto Content and compare the four formats and review scores.",
        "Apply one revision request, approve the final suite, and download JSON.",
        "Run the offline evaluator to demonstrate an Overall: PASS result.",
    ])

    add_heading(doc, "9. Conclusion")
    add_body(doc, "AI at Your Doorstep satisfies the assignment’s MVP goal: it accepts a campaign input, discovers or ingests source material, uses RAG to retrieve grounded context, generates a complete GTM suite, and runs automated critique and revision. The implementation is deliberately human-approved at every consequential boundary and provides evidence-linked structured outputs rather than untraceable prose.")
    add_body(doc, "The strongest technical improvement was moving chunking from a documented assumption to a corpus-tied, retrieval-tested design. The resulting 3,600 / 450 configuration materially improved strict passage recall on the challenge set. Together with 71 passing tests and a successful exported content suite, this provides reproducible evidence that the project is ready as a local assignment and portfolio MVP.")

    add_heading(doc, "Appendix A — Reproducible commands", page_break_before=True)
    commands = [
        ("Create environment", "python3.11 -m venv .venv\nsource .venv/bin/activate\npip install -r requirements-agent.txt"),
        ("Start application", "streamlit run app.py"),
        ("Run tests", "python -m unittest discover -s tests"),
        ("Profile corpus", "python scripts/profile_chunking.py data/corpus"),
        ("Run baseline retrieval", "python scripts/evaluate_retrieval.py --top-k 5"),
        ("Compare chunking", "python scripts/compare_chunking_retrieval.py --repetitions 3"),
        ("Evaluate exported suite", "python scripts/evaluate_content_suite.py ai-week-in-review-content-suite.json"),
    ]
    for label, command in commands:
        p = doc.add_paragraph()
        style_paragraph(p, after=2, before=6)
        set_run(p.add_run(label), bold=True, color=BLUE)
        box = doc.add_table(rows=1, cols=1)
        set_table_geometry(box, [TABLE_WIDTH])
        shade(box.cell(0, 0), LIGHT_GRAY)
        cp = box.cell(0, 0).paragraphs[0]
        style_paragraph(cp, after=0, line=1.0)
        set_run(cp.add_run(command), size=9.2, font="Courier New", color=DARK)

    add_heading(doc, "Appendix B — Key module map")
    add_table(doc, ["Path", "Purpose"], [
        ["app.py", "Streamlit UI, session state, approval gates, and exports"],
        ["agents/main_agent.py", "Deterministic coordinator interface"],
        ["agents/deep_runtime.py", "Deep Agent construction and call limits"],
        ["agents/tool_adapters.py", "Approval-aware tools and skipped-step fallback"],
        ["agents/content_writer.py", "Four-format structured generation and validation"],
        ["agents/reviewer.py", "Structured critique, revision, and grounding validation"],
        ["research/workflow.py", "Discovery, approved ingestion, retrieval, and brief workflow"],
        ["research/pinecone_store.py", "Vector index management, metadata, queries, and filters"],
        ["research/retrieval.py", "Editorial multi-query retrieval and deduplication"],
        ["ingestion/chunker.py", "Paragraph-aware 3,600 / 450 chunking and stable IDs"],
        ["evaluations/", "Content, bundle, chunking, baseline, and challenge evaluations"],
    ], [3000, 6360])
    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build()
    document.core_properties.title = "AI at Your Doorstep — GTM Agent Project Report"
    document.core_properties.subject = "Assignment submission"
    document.core_properties.author = "Ragavi Suriyan"
    document.core_properties.keywords = "GTM agent, RAG, Pinecone, OpenAI, Deep Agent"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
