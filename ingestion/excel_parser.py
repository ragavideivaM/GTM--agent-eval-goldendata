"""Extract readable row-oriented text from Excel workbooks."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class ParsedExcel:
    filename: str
    text: str
    sheet_names: tuple[str, ...]
    row_count: int


def extract_excel(workbook_bytes: bytes, filename: str = "uploaded.xlsx") -> ParsedExcel:
    """Convert non-empty workbook cells into labeled, searchable text rows."""
    if not workbook_bytes:
        raise ValueError("Excel file cannot be empty.")
    try:
        workbook = load_workbook(
            BytesIO(workbook_bytes), read_only=True, data_only=True
        )
    except Exception as exc:
        raise ValueError("The uploaded file is not a readable .xlsx workbook.") from exc

    sections: list[str] = []
    row_count = 0
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row_number, cells in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() for value in cells if value is not None and str(value).strip()]
                if values:
                    rows.append(f"Row {row_number}: " + " | ".join(values))
                    row_count += 1
            if rows:
                sections.append(f"[Sheet: {worksheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()

    if not sections:
        raise ValueError("The Excel workbook contains no readable values.")
    return ParsedExcel(
        filename=Path(filename).name,
        text="\n\n".join(sections),
        sheet_names=tuple(workbook.sheetnames),
        row_count=row_count,
    )
